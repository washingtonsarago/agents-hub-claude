"""Generate a single-slide PPTX summarizing the agents-hub-claude project and its benefits."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Brand colors
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x2E, 0x86, 0xDE)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
GREEN = RGBColor(0x27, 0xAE, 0x60)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

# --- Background ---
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = LIGHT_BG
bg.line.fill.background()

# --- Top navy header bar ---
header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.4))
header.fill.solid()
header.fill.fore_color.rgb = NAVY
header.line.fill.background()

# Title
title_tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.5), Inches(0.75))
tf = title_tb.text_frame
tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run = p.add_run()
run.text = "agents-hub-claude"
run.font.name = "Calibri"
run.font.size = Pt(36)
run.font.bold = True
run.font.color.rgb = WHITE

# Subtitle
subtitle_tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.85), Inches(12.5), Inches(0.45))
tf = subtitle_tb.text_frame
tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
run = p.add_run()
run.text = "Registry centralizado de agents e commands do Claude Code — auto-atualizado em toda máquina da EMS-NCTECH"
run.font.name = "Calibri"
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0xCF, 0xDC, 0xEB)

# --- Section: "O que é?" (left column) ---
what_title = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(5.8), Inches(0.5))
tf = what_title.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "O que é"
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = ACCENT
run.font.name = "Calibri"

what_body = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(5.8), Inches(4.8))
tf = what_body.text_frame
tf.word_wrap = True

items = [
    ("Hub único", "Um repo Git (EMS-NCTECH/agents-hub-claude) com 11 agents + 8 slash commands versionados."),
    ("CLI própria (ahc)", "Node zero-deps que sincroniza tudo via git clone autenticado, compatível com repos INTERNAL."),
    ("Auto-update transparente", "Hook SessionStart do Claude Code roda ahc sync a cada abertura — sem intervenção manual."),
    ("Integridade garantida", "manifest.json com sha256 por item + lock file local; só baixa o que realmente mudou."),
]

for i, (head, body) in enumerate(items):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.space_after = Pt(10)
    r1 = p.add_run()
    r1.text = f"▪ {head}  "
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = TEXT
    r1.font.name = "Calibri"
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(12)
    r2.font.color.rgb = MUTED
    r2.font.name = "Calibri"

# --- Divider ---
div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.65), Inches(1.8), Emu(12700), Inches(5.3))
div.fill.solid()
div.fill.fore_color.rgb = RGBColor(0xD9, 0xE1, 0xEB)
div.line.fill.background()

# --- Section: "Benefícios" (right column) ---
ben_title = slide.shapes.add_textbox(Inches(6.9), Inches(1.7), Inches(6.0), Inches(0.5))
tf = ben_title.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "Benefícios"
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = ACCENT
run.font.name = "Calibri"

benefits = [
    ("Padronização", "Todo dev da EMS usa os mesmos agents e commands — sem divergência entre máquinas."),
    ("Velocidade", "Novo dev produtivo em minutos: um script de install entrega 19 ferramentas prontas."),
    ("Governança", "Mudanças passam por PR + review. Versionamento semântico e rollback via ahc pin."),
    ("Segurança", "Todos os agents seguem OWASP Top 10, SoC/DRY/KISS/YAGNI/SOLID; sem MediatR (pago)."),
    ("Colaboração", "Agents têm Collaboration Protocol padronizado — delegam entre si de forma consistente."),
    ("Evolução contínua", "Commit → push → toda a org recebe no próximo SessionStart. Zero fricção."),
]

ben_body = slide.shapes.add_textbox(Inches(6.9), Inches(2.2), Inches(6.0), Inches(4.9))
tf = ben_body.text_frame
tf.word_wrap = True

for i, (head, body) in enumerate(benefits):
    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
    p.space_after = Pt(6)
    r1 = p.add_run()
    r1.text = "✓ "
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = GREEN
    r1.font.name = "Calibri"
    r2 = p.add_run()
    r2.text = f"{head}  "
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = TEXT
    r2.font.name = "Calibri"
    r3 = p.add_run()
    r3.text = body
    r3.font.size = Pt(11)
    r3.font.color.rgb = MUTED
    r3.font.name = "Calibri"

# --- Footer stats band ---
footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), prs.slide_width, Inches(0.6))
footer.fill.solid()
footer.fill.fore_color.rgb = NAVY
footer.line.fill.background()

stats = slide.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.5), Inches(0.5))
tf = stats.text_frame
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER

stats_line = [
    ("11", " agents   •   "),
    ("8", " commands   •   "),
    ("git-clone auth", "   •   "),
    ("SessionStart hook", "   •   "),
    ("sha256", " integrity   •   "),
    ("github.com/EMS-NCTECH/agents-hub-claude", ""),
]
for val, sep in stats_line:
    r = p.add_run()
    r.text = val
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = "Calibri"
    if sep:
        r2 = p.add_run()
        r2.text = sep
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(0xA8, 0xBB, 0xD0)
        r2.font.name = "Calibri"

out = "docs/agents-hub-claude.pptx"
prs.save(out)
print(f"saved: {out}")
